"""Import the sponsor's CAC workbook without inventing or verifying source facts."""
import datetime as dt
import hashlib
import re
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from django.utils import timezone

from .models import Conference, ConferenceEdition, InformationStatus, VerificationStatus
from .management.commands.import_legacy_excel import parse_format

MONTHS = {name[:3].upper(): number for number, name in enumerate(
    ['January','February','March','April','May','June','July','August','September','October','November','December'],1)}
TRACKING = {'gclid','gbraid','wbraid','fbclid','gad_source','gad_campaignid'}


def compact(value):
    return re.sub(r'\s+', ' ', str('' if value is None else value)).strip()


def canonical_url(value):
    value = compact(value)
    if not value:
        return None
    if not re.match(r'^https?://', value, re.I):
        value = 'https://' + value
    parts = urlsplit(value)
    query = [(k,v) for k,v in parse_qsl(parts.query,keep_blank_values=True)
             if not k.lower().startswith('utm_') and k.lower() not in TRACKING]
    value = urlunsplit((parts.scheme.lower(),parts.netloc.lower(),parts.path,urlencode(query),parts.fragment))
    try:
        URLValidator(schemes=['http','https'])(value)
    except ValidationError:
        return None
    # Preserve the full original in source metadata; never silently truncate a link.
    return value if len(value) <= 200 else None


def website_identity(value):
    value = canonical_url(value)
    if not value:
        return None
    p = urlsplit(value)
    return (p.netloc.lower().removeprefix('www.'),p.path.rstrip('/'),p.query)


def listed_dates(value, context_year, name='', url=''):
    warnings = []
    start = end = None
    if isinstance(value,dt.datetime): value = value.date()
    if isinstance(value,dt.date):
        return value,None,value.year,'date cell',warnings
    raw = compact(value).upper().replace('â€“','-').replace('â€”','-')
    match = re.search(r'\b(20\d{2})\b', raw)
    short = re.search(r'\s(\d{2})$',raw)
    explicit = int(match.group(1)) if match else 2000+int(short.group(1)) if short else None
    named = re.search(r'\b(20\d{2})\b', name)
    path = urlsplit(url).path
    url_year = re.search(r'(?<![A-Za-z0-9])(20\d{2})(?!\d)',path)
    shorthand = re.search(r'/atc(\d{2})(?:/|$)',path,re.I)
    hinted = int(named.group(1)) if named else int(url_year.group(1)) if url_year else 2000+int(shorthand.group(1)) if shorthand else None
    year = explicit or hinted or context_year
    basis = 'date cell' if explicit else 'conference name' if named else 'website URL' if hinted else 'workbook heading'
    if not year:
        raise ValueError('No edition year can be established from the supplied source.')
    if not explicit:
        warnings.append(f'Edition year {year} inferred from {basis}; verify the event year.')
    if year != context_year:
        warnings.append(f'Listed edition is {year}; the workbook heading is {context_year}.')
    if raw:
        patterns = [
            (r'^(\d{1,2})\s*-\s*(\d{1,2})\s+([A-Z]+)(?:\s+\d{2,4})?$', 'same'),
            (r'^(\d{1,2})\s+([A-Z]+)\s*-\s*(\d{1,2})\s+([A-Z]+)(?:\s+\d{2,4})?$', 'cross'),
            (r'^(\d{1,2})\s+([A-Z]+)(?:\s+\d{2,4})?$', 'single'),
        ]
        for pattern,kind in patterns:
            m = re.fullmatch(pattern,raw)
            if not m: continue
            try:
                if kind == 'same':
                    month=MONTHS[m[3][:3]]
                    start,end=dt.date(year,month,int(m[1])),dt.date(year,month,int(m[2]))
                elif kind == 'cross':
                    start=dt.date(year,MONTHS[m[2][:3]],int(m[1]))
                    month=MONTHS[m[4][:3]]
                    end=dt.date(year+(month<start.month),month,int(m[3]))
                else:
                    start=dt.date(year,MONTHS[m[2][:3]],int(m[1]))
                if end and end < start: raise ValueError()
            except (KeyError,ValueError):
                start=end=None
            break
        if not start:
            warnings.append(f'Date text could not be normalized: {compact(value)}. Original retained.')
        elif end and (end-start).days > 10:
            warnings.append('The listed date range exceeds ten days; verify the source dates.')
    else:
        warnings.append('Event dates are missing from the spreadsheet.')
    return start,end,year,basis,warnings


def workbook_records(path):
    import openpyxl
    path=Path(path)
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    wb=openpyxl.load_workbook(path,data_only=True)
    records=[]
    try:
        for ws in wb:
            header_row=None
            for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,20)):
                labels=[compact(c.value).upper() for c in row]
                if 'NAME' in labels and 'DATE' in labels and 'LINK' in labels:
                    header_row=row[0].row
                    break
            if header_row is None: continue
            heading=' '.join(compact(c.value) for row in ws.iter_rows(max_row=header_row-1) for c in row)
            match=re.search(r'\b(20\d{2})\b',heading)
            if not match or 'CAC' not in heading.upper():
                raise ValueError('This command requires a CAC heading with an explicit listing year.')
            context_year=int(match.group(1))
            headers={compact(c.value).upper():c.column for c in ws[header_row] if compact(c.value)}
            if len(headers)!=sum(bool(compact(c.value)) for c in ws[header_row]):
                raise ValueError('Duplicate spreadsheet column labels need review before import.')
            for row in ws.iter_rows(min_row=header_row+1):
                def cell(label):
                    col=headers.get(label)
                    return ws.cell(row[0].row,col) if col else None
                def value(label):
                    c=cell(label)
                    return c.value if c else None
                name=compact(value('NAME'))
                if not name: continue
                if len(name)>255: raise ValueError(f'Row {row[0].row}: conference name exceeds 255 characters.')
                c=cell('LINK')
                original_url=(c.hyperlink.target if c and c.hyperlink else value('LINK')) or ''
                url=canonical_url(original_url)
                start,end,year,basis,warnings=listed_dates(value('DATE'),context_year,name,str(original_url))
                if original_url and not url:
                    warnings.append('Website link is invalid or exceeds 200 characters after removing tracking; original retained.')
                raw={}
                for label,col in headers.items():
                    v=ws.cell(row[0].row,col).value
                    raw[label]=v.isoformat() if isinstance(v,(dt.datetime,dt.date)) else v
                cost=value('COST')
                if cost is not None and compact(cost):
                    warnings.append('Cost retained as source text; confirm amount and currency before creating a registration price.')
                if start and isinstance(value('DATE'),(dt.date,dt.datetime)):
                    warnings.append('Only one date is listed; no end date was assumed.')
                records.append({'name':name,'url':url,'year':year,'start':start,'end':end,
                    'location':compact(value('LOCATION')) or None,'format':parse_format(value('VIRTUAL/IN PERSON')),
                    'category':compact(value('CATEGORY')),'source':{
                        'id':f'{digest}:{ws.title}:{row[0].row}', 'filename':path.name,'sha256':digest,
                        'sheet':ws.title,'row':row[0].row,'heading':heading,'raw_values':raw,
                        'original_url':str(original_url),'listed_date':raw.get('DATE'),
                        'focus_area':compact(value('FOCUS AREA')),'category':compact(value('CATEGORY')),
                        'notes':str(value('NOTES') or '').strip(),
                        'other_dates':str(value('OTHER CONFERENCE DATES') or '').strip(),
                        'year_basis':basis,'warnings':warnings,
                    }})
    finally:
        wb.close()
    if not records: raise ValueError('No records found under a NAME / DATE / LINK header.')
    return records


def import_workbook(path, *, dry_run=False):
    records=workbook_records(path)
    report={'source':Path(path).name,'rows':len(records),'conferences_created':0,'editions_created':0,
            'already_imported':0,'existing_editions_preserved':0,'prices_created':0,'records':[]}
    with transaction.atomic():
        conferences=list(Conference.objects.all())
        for record in records:
            source=record['source']
            identity=website_identity(record['url'])
            matches=[c for c in conferences if (identity and website_identity(c.website)==identity)
                     or compact(c.name).casefold()==record['name'].casefold()]
            if len(matches)>1:
                raise ValueError(f"Row {source['row']}: multiple existing conferences match; resolve before import.")
            c=matches[0] if matches else None
            if c is None:
                c=Conference.objects.create(name=record['name'],website=record['url'],cac_relevant=True,
                    themes=[record['category']] if record['category'] else [])
                conferences.append(c)
                report['conferences_created']+=1
            elif not c.cac_relevant:
                c.cac_relevant=True
                c.save(update_fields=['cac_relevant'])
            edition,created=ConferenceEdition.objects.get_or_create(conference=c,year=record['year'],defaults={
                'start_date':record['start'],'end_date':record['end'],'location':record['location'],
                'attendance_format':record['format'],'source_url':record['url'],
                'information_status': InformationStatus.NEEDS_VERIFICATION if record['start'] and record['location'] else InformationStatus.PARTIAL_INFORMATION,
                'verification_status':VerificationStatus.NEEDS_REVIEW,
            })
            sources=edition.field_sources.get('legacy_imports',[])
            if any(s.get('id')==source['id'] for s in sources):
                report['already_imported']+=1
                action='already imported'
            else:
                snapshot=dict(source,imported_at=timezone.now().isoformat())
                edition.field_sources={**edition.field_sources,'legacy_imports':[*sources,snapshot]}
                edition.save(update_fields=['field_sources'])
                if created:
                    report['editions_created']+=1
                    action='created unverified edition'
                else:
                    report['existing_editions_preserved']+=1
                    action='attached source; existing edition retained'
            report['records'].append({'row':source['row'],'source_name':record['name'],'catalog_name':c.name,
                'conference_id':c.pk,'edition_id':edition.pk,'year':record['year'],'year_basis':source['year_basis'],
                'start_date':record['start'].isoformat() if record['start'] else None,
                'end_date':record['end'].isoformat() if record['end'] else None,'action':action,'warnings':source['warnings']})
        if dry_run: transaction.set_rollback(True)
    report['dry_run']=dry_run
    return report
