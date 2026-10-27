import csv
import datetime
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from conferences.models import (
    AttendanceFormat,
    Conference,
    ConferenceEdition,
    Deadline,
    InformationStatus,
    RegistrationPrice,
    VerificationStatus,
)

ALIASES = {
    "name": ["conference", "conferencename", "name", "event", "eventname", "title"],
    "organizer": ["organizer", "organization", "host", "sponsor", "society"],
    "website": ["website", "url", "link", "conferencewebsite", "web", "site"],
    "year": ["year", "edition", "conferenceyear"],
    "start_date": ["startdate", "start", "begins", "datestart", "fromdate"],
    "end_date": ["enddate", "end", "ends", "dateend", "todate"],
    "dates": ["date", "dates", "conferencedates", "eventdates", "when"],
    "location": ["location", "city", "venue", "place", "citystate", "where"],
    "format": ["format", "attendanceformat", "type", "mode", "inpersonvirtual", "modality"],
    "themes": ["themes", "theme", "topics", "topic", "category", "categories", "focus", "focusarea"],
    "audience": ["audience", "targetaudience", "attendees"],
    "center": ["center", "centers", "caccdsa", "relevantcenter", "department", "unit"],
    "cac": ["cac", "cacrelevant", "appliedcomputing"],
    "cdsa": ["cdsa", "cdsarelevant", "datascience"],
    "cost": ["cost", "registration", "registrationcost", "registrationfee", "fee", "price", "regfee"],
    "registration_deadline": ["registrationdeadline", "regdeadline", "deadlinetoregister"],
    "early_registration_deadline": ["earlyregistrationdeadline", "earlybirddeadline", "earlybird", "earlyregdeadline"],
    "submission_deadline": ["submissiondeadline", "cfpdeadline", "callforpapers", "paperdeadline", "abstractdeadline", "cfp"],
    "deadline": ["deadline", "deadlines", "keydeadline"],
}

DEADLINE_COLUMNS = {
    "registration_deadline": "registration",
    "early_registration_deadline": "early_registration",
    "submission_deadline": "abstract_submission",
    "deadline": "other",
}

DATE_FORMATS = [
    "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%B %d, %Y", "%b %d, %Y",
    "%B %d %Y", "%b %d %Y", "%d %B %Y", "%d %b %Y", "%Y/%m/%d",
]

MONTHS = {m.lower(): i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], start=1)}
MONTHS.update({k[:3]: v for k, v in list(MONTHS.items())})
MONTHS["sept"] = 9


def squash(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_date(value):
    if blank(value):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", str(value).strip())
    text = re.sub(r"\s+", " ", text.replace(".", ""))
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_range(value):
    if blank(value):
        return None, None
    if isinstance(value, (datetime.date, datetime.datetime)):
        d = parse_date(value)
        return d, d
    text = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", str(value).strip())
    text = text.replace("–", "-").replace("—", "-").replace(" to ", "-")
    single = parse_date(text)
    if single:
        return single, single

    m = re.match(r"^([A-Za-z]+)\.?\s+(\d{1,2})\s*-\s*(\d{1,2}),?\s+(\d{4})$", text)
    if m and m.group(1).lower() in MONTHS:
        month, year = MONTHS[m.group(1).lower()], int(m.group(4))
        try:
            return datetime.date(year, month, int(m.group(2))), datetime.date(year, month, int(m.group(3)))
        except ValueError:
            return None, None

    m = re.match(r"^([A-Za-z]+)\.?\s+(\d{1,2})\s*-\s*([A-Za-z]+)\.?\s+(\d{1,2}),?\s+(\d{4})$", text)
    if m and m.group(1).lower() in MONTHS and m.group(3).lower() in MONTHS:
        year = int(m.group(5))
        try:
            return (datetime.date(year, MONTHS[m.group(1).lower()], int(m.group(2))),
                    datetime.date(year, MONTHS[m.group(3).lower()], int(m.group(4))))
        except ValueError:
            return None, None

    parts = [p.strip() for p in text.split("-") if p.strip()]
    if len(parts) == 2:
        start, end = parse_date(parts[0]), parse_date(parts[1])
        if start and end:
            return start, end
    return None, None


def parse_year(value):
    if blank(value):
        return None
    if isinstance(value, (int, float)) and 1900 < int(value) < 2100:
        return int(value)
    m = re.search(r"\b(19|20)\d{2}\b", str(value))
    return int(m.group(0)) if m else None


def parse_money(value):
    if blank(value):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    m = re.search(r"\d[\d,]*(\.\d+)?", str(value))
    if not m:
        return None
    try:
        return Decimal(m.group(0).replace(",", "")).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def parse_yes(value):
    if blank(value):
        return False
    if isinstance(value, bool):
        return value
    return squash(value) in {"y", "yes", "x", "true", "1", "relevant"}


def parse_list(value):
    if blank(value):
        return []
    return [p.strip() for p in re.split(r"[,;/\n]", str(value)) if p.strip()]


def parse_format(value):
    text = squash(value)
    if not text:
        return None
    if "hybrid" in text or ("person" in text and "virtual" in text):
        return AttendanceFormat.HYBRID
    if "virtual" in text or "online" in text or "remote" in text:
        return AttendanceFormat.VIRTUAL
    if "person" in text or "onsite" in text:
        return AttendanceFormat.IN_PERSON
    return None


def clean_url(value):
    if blank(value):
        return None
    text = str(value).strip()
    if not re.match(r"^https?://", text, re.I):
        if "." not in text or " " in text:
            return None
        text = "https://" + text
    return text[:200]


def read_rows(path, sheet, header_row):
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    else:
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl isn't installed. Run: pip install openpyxl")
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet and sheet not in wb.sheetnames:
            raise CommandError(f"No sheet named {sheet!r}. Sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    if len(rows) < header_row:
        raise CommandError("The file doesn't have a header row where expected.")
    headers = [str(h).strip() if h is not None else "" for h in rows[header_row - 1]]
    return headers, rows[header_row:]


def map_headers(headers):
    lookup = {}
    for key, names in ALIASES.items():
        for n in names:
            lookup.setdefault(n, key)
    mapping, unmapped = {}, []
    for index, header in enumerate(headers):
        key = lookup.get(squash(header))
        if key and key not in mapping:
            mapping[key] = index
        elif header:
            unmapped.append(header)
    return mapping, unmapped


class Command(BaseCommand):
    help = "Import conferences from the legacy CCSE spreadsheet (.xlsx or .csv)"

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--sheet", help="Sheet name (defaults to the first sheet)")
        parser.add_argument("--header-row", type=int, default=1)
        parser.add_argument("--default-year", type=int, help="Year to use when a row has no year or date")
        parser.add_argument("--dry-run", action="store_true", help="Show what would happen without saving")

    def handle(self, *args, **opts):
        path = Path(opts["path"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        headers, rows = read_rows(path, opts["sheet"], opts["header_row"])
        mapping, unmapped = map_headers(headers)
        if "name" not in mapping:
            raise CommandError(
                "Couldn't find a conference name column. Headers found: " + ", ".join(h for h in headers if h)
            )

        self.stdout.write("Column mapping:")
        for key, index in mapping.items():
            self.stdout.write(f"  {headers[index]!r} -> {key}")
        if unmapped:
            self.stdout.write(self.style.WARNING("Not imported (no matching field): " + ", ".join(unmapped)))

        stats = {"conferences_created": 0, "editions_created": 0, "editions_skipped": 0,
                 "deadlines": 0, "prices": 0, "rows_skipped": 0}
        problems = []

        with transaction.atomic():
            for line, row in enumerate(rows, start=opts["header_row"] + 1):
                get = lambda key: row[mapping[key]] if key in mapping and mapping[key] < len(row) else None
                name = get("name")
                if blank(name):
                    continue
                name = re.sub(r"\s+", " ", str(name)).strip()[:255]
                self._import_row(line, name, get, opts["default_year"], stats, problems)

            if opts["dry_run"]:
                transaction.set_rollback(True)

        for p in problems:
            self.stdout.write(self.style.WARNING(p))
        label = "Dry run, nothing saved" if opts["dry_run"] else "Import finished"
        self.stdout.write(self.style.SUCCESS(
            f"{label}: {stats['conferences_created']} conferences and {stats['editions_created']} editions created, "
            f"{stats['editions_skipped']} editions already existed, {stats['deadlines']} deadlines, "
            f"{stats['prices']} prices, {stats['rows_skipped']} rows skipped."
        ))

    def _import_row(self, line, name, get, default_year, stats, problems):
        start, end = parse_date(get("start_date")), parse_date(get("end_date"))
        if not start and not blank(get("dates")):
            start, end = parse_range(get("dates"))
            if not start:
                problems.append(f"Row {line}: couldn't read dates {get('dates')!r}")
        if start and not end:
            end = start

        year = parse_year(get("year")) or (start.year if start else None) or default_year
        if not year:
            stats["rows_skipped"] += 1
            problems.append(f"Row {line}: {name!r} has no year or date, skipped (use --default-year)")
            return

        center = squash(get("center"))
        cac = parse_yes(get("cac")) or "cac" in center or "ccse" in center
        cdsa = parse_yes(get("cdsa")) or "cdsa" in center or "ccse" in center
        website = clean_url(get("website"))

        conference = Conference.objects.filter(name__iexact=name).first()
        if conference is None:
            conference = Conference.objects.create(
                name=name,
                organizer=(str(get("organizer")).strip()[:255] if not blank(get("organizer")) else None),
                website=website,
                themes=parse_list(get("themes")),
                target_audience=parse_list(get("audience")),
                cac_relevant=cac,
                cdsa_relevant=cdsa,
            )
            stats["conferences_created"] += 1
        else:
            changed = False
            if cac and not conference.cac_relevant:
                conference.cac_relevant = changed = True
            if cdsa and not conference.cdsa_relevant:
                conference.cdsa_relevant = changed = True
            if website and not conference.website:
                conference.website, changed = website, True
            if changed:
                conference.save()

        if ConferenceEdition.objects.filter(conference=conference, year=year).exists():
            stats["editions_skipped"] += 1
            problems.append(f"Row {line}: {name} {year} already exists, left as is")
            return

        location = get("location")
        format_value = get("format")
        attendance = parse_format(format_value)
        if not blank(format_value) and attendance is None:
            problems.append(f"Row {line}: couldn't read format {format_value!r}")

        has_details = bool(start and location and not blank(location))
        edition = ConferenceEdition.objects.create(
            conference=conference,
            year=year,
            start_date=start,
            end_date=end,
            location=(str(location).strip()[:255] if not blank(location) else None),
            attendance_format=attendance,
            information_status=(InformationStatus.NEEDS_VERIFICATION if has_details
                                else InformationStatus.PARTIAL_INFORMATION),
            verification_status=VerificationStatus.NEEDS_REVIEW,
            source_url=website,
        )
        stats["editions_created"] += 1

        for column, deadline_type in DEADLINE_COLUMNS.items():
            raw = get(column)
            if blank(raw):
                continue
            when = parse_date(raw)
            if when is None:
                problems.append(f"Row {line}: couldn't read {column.replace('_', ' ')} {raw!r}")
                continue
            Deadline.objects.create(
                edition=edition, type=deadline_type, date=when,
                evidence="Imported from legacy spreadsheet",
            )
            stats["deadlines"] += 1

        raw_cost = get("cost")
        amount = parse_money(raw_cost)
        if amount is not None:
            RegistrationPrice.objects.create(
                edition=edition, category="General", rate_type="legacy_listed", amount=amount,
                evidence=f"Legacy spreadsheet value: {raw_cost}",
            )
            stats["prices"] += 1
        elif not blank(raw_cost):
            problems.append(f"Row {line}: couldn't read cost {raw_cost!r}")
